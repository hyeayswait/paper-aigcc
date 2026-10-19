# -*- coding: utf-8 -*-
"""非正文行识别的回归测试（真实论文压出的 bug）。"""
from __future__ import annotations

import os
import tempfile
import unittest

from docx import Document as DocxDocument

from paper_aigc.core.textnorm import is_furniture_like
from paper_aigc.io.docx_reader import read_docx


class FurnitureTest(unittest.TestCase):
    def test_prose_starting_with_table_number_is_kept(self) -> None:
        text = "表4里，营业总收入从1234.56亿元涨到2345.67亿元，营业利润保持两位数增长。"
        self.assertFalse(is_furniture_like(text))

    def test_caption_is_removed(self) -> None:
        self.assertTrue(is_furniture_like("表1 某公司2023-2025年盈利能力指标"))
        self.assertTrue(is_furniture_like("图2 存货周转率变化趋势"))

    def test_keyword_and_note_are_removed(self) -> None:
        self.assertTrue(is_furniture_like("关键词：某公司；财务绩效"))
        self.assertTrue(is_furniture_like("注：数据来源于公司年报"))

    def test_author_line_is_removed(self) -> None:
        self.assertTrue(is_furniture_like("某某学院会计2401班，指导教师：张某某"))

    def test_reader_keeps_body_paragraph_starting_with_table_ref(self) -> None:
        temp_dir = tempfile.mkdtemp()
        path = os.path.join(temp_dir, "case.docx")
        doc = DocxDocument()
        heading = doc.add_paragraph()
        heading.style = doc.styles["Heading 1"]
        heading.add_run("一、分析")
        body = doc.add_paragraph(
            "表4里，营业总收入从1234.56亿元涨到2345.67亿元，营业利润保持两位数增长。"
        )
        caption = doc.add_paragraph("表4 成长能力指标")
        doc.save(path)
        parsed = read_docx(path)
        texts = [parsed.paragraphs[i].text for i in parsed.body_indices]
        self.assertIn(body.text, texts)
        self.assertNotIn(caption.text, texts)
        self.assertEqual(len(parsed.body_indices), 1)


if __name__ == "__main__":
    unittest.main()
