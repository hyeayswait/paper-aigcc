"""T2 单元测试：docx 解析 / 回写 / 报告渲染。

所有临时 docx 与报告都写到 tempfile.mkdtemp() 下，不往仓库里丢文件；
不依赖网络、API Key，也不依赖 T1 的 features/scoring 模块。
"""
from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import unittest

import docx
from docx.shared import Pt

from paper_aigc.core.types import (
    DISCLAIMER_ZH,
    VERSION,
    Document,
    Paragraph,
    ParagraphScore,
    RewriteResult,
    ScoreReport,
)
from paper_aigc.io.docx_reader import read_docx
from paper_aigc.io.docx_writer import write_docx
from paper_aigc.io.report import render_html, render_markdown

HEADING_TEXT = "第一章 绪论"
BODY1 = "本研究围绕论文的文体特征展开分析，并给出可复核的指标定义。"
BODY2 = "在方法上，本文采用分段统计与加权汇总的方式构造指数。"
SHORT_TEXT = "结论如下。"
REF_TEXT = "[1] 张三. 论文题名[J]. 期刊名称, 2020, 12(3): 45-52."
XSS_TEXT = '<script>alert("xss")</script> 段落包含 & 与 < > 符号。'
LOW_TEXT = "这是低风险段落原文。"


class _TempDirMixin:
    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="paperaigc_io_")

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)


def _make_sample_docx(path: str) -> None:
    document = docx.Document()
    heading = document.add_paragraph(HEADING_TEXT)
    heading.style = document.styles["Heading 1"]
    document.add_paragraph(BODY1)
    document.add_paragraph(BODY2)
    document.add_paragraph(SHORT_TEXT)
    document.add_paragraph(REF_TEXT)
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.save(path)


def _hash_file(path: str) -> str:
    with open(path, "rb") as handle:
        return hashlib.sha256(handle.read()).hexdigest()


def _sample_report() -> ScoreReport:
    document = Document(
        path="测试论文.docx",
        title="测试论文",
        paragraphs=[
            Paragraph(index=0, text=LOW_TEXT),
            Paragraph(index=1, text=XSS_TEXT),
        ],
        body_indices=[0, 1],
        total_words=200,
    )
    paragraphs = [
        ParagraphScore(
            index=0,
            text=LOW_TEXT,
            score=22.5,
            level="低",
            top_signals=[("short_sentence_ratio", "短句占比", 4.2)],
            suggestion="保持长短句交替。",
        ),
        ParagraphScore(
            index=1,
            text=XSS_TEXT,
            score=88.0,
            level="高",
            top_signals=[
                ("boilerplate_density", "套话密度", 21.3),
                ("connective_density", "连接词密度", 12.0),
            ],
            suggestion="删掉套话 <b>并拆句</b>。",
        ),
    ]
    return ScoreReport(
        document=document,
        paragraphs=paragraphs,
        overall=61.3,
        level="高",
        level_counts={"低": 1, "中": 0, "偏高": 0, "高": 1},
        meta={"version": VERSION, "feature_count": 17, "engine": "stylometric-heuristic"},
    )


class DocxReaderTest(_TempDirMixin, unittest.TestCase):
    def test_read_docx_structure(self) -> None:
        path = os.path.join(self.tmp, "sample.docx")
        _make_sample_docx(path)
        document = read_docx(path)

        self.assertEqual(document.title, HEADING_TEXT)
        self.assertEqual(len(document.paragraphs), 5)
        self.assertEqual([p.index for p in document.paragraphs], [0, 1, 2, 3, 4])
        self.assertTrue(document.paragraphs[0].is_heading)
        self.assertTrue(document.paragraphs[4].is_reference)
        self.assertFalse(document.paragraphs[3].is_heading)
        self.assertLess(document.paragraphs[3].char_count, 10)
        self.assertEqual(document.body_indices, [1, 2])
        self.assertNotIn(3, document.body_indices)
        self.assertNotIn(4, document.body_indices)
        self.assertEqual(
            document.total_words,
            document.paragraphs[1].char_count + document.paragraphs[2].char_count,
        )

    def test_read_docx_title_fallback(self) -> None:
        path = os.path.join(self.tmp, "no_heading.docx")
        document = docx.Document()
        document.add_paragraph(BODY1)
        document.save(path)
        read = read_docx(path)
        self.assertEqual(read.title, "no_heading")
        self.assertEqual(read.body_indices, [0])

    def test_read_docx_empty_document(self) -> None:
        path = os.path.join(self.tmp, "empty.docx")
        document = docx.Document()
        document.add_paragraph("")
        document.save(path)
        read = read_docx(path)
        self.assertEqual(read.paragraphs, [])
        self.assertEqual(read.body_indices, [])
        self.assertEqual(read.total_words, 0)

    def test_read_docx_ignores_tables(self) -> None:
        path = os.path.join(self.tmp, "table.docx")
        document = docx.Document()
        document.add_paragraph(BODY1)
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格里的文字不应被收录"
        document.save(path)
        read = read_docx(path)
        self.assertEqual(len(read.paragraphs), 1)
        self.assertNotIn("表格里的文字不应被收录", [p.text for p in read.paragraphs])

    def test_read_docx_normalizes_whitespace(self) -> None:
        path = os.path.join(self.tmp, "spaces.docx")
        document = docx.Document()
        document.add_paragraph("  甲乙    丙丁  ")
        document.save(path)
        read = read_docx(path)
        self.assertEqual(read.paragraphs[0].text, "甲乙 丙丁")

    def test_read_docx_missing_file_raises_value_error(self) -> None:
        missing = os.path.join(self.tmp, "not_here.docx")
        with self.assertRaises(ValueError) as ctx:
            read_docx(missing)
        self.assertIn("不存在", str(ctx.exception))

    def test_read_docx_wrong_extension_raises_value_error(self) -> None:
        wrong = os.path.join(self.tmp, "paper.txt")
        with open(wrong, "w", encoding="utf-8") as handle:
            handle.write("hello")
        with self.assertRaises(ValueError) as ctx:
            read_docx(wrong)
        self.assertIn(".docx", str(ctx.exception))

    def test_read_docx_broken_file_raises_value_error(self) -> None:
        broken = os.path.join(self.tmp, "broken.docx")
        with open(broken, "wb") as handle:
            handle.write(b"this is not a real docx package")
        with self.assertRaises(ValueError) as ctx:
            read_docx(broken)
        self.assertIn("无法解析", str(ctx.exception))


class DocxWriterTest(_TempDirMixin, unittest.TestCase):
    def test_write_docx_replaces_only_target_paragraph(self) -> None:
        source = os.path.join(self.tmp, "source.docx")
        _make_sample_docx(source)
        before_hash = _hash_file(source)
        out_path = os.path.join(self.tmp, "out.docx")
        replacement = "改写后的第一段：句式更自然，也更简短。"

        returned = write_docx(source, {1: replacement}, out_path)

        self.assertEqual(returned, out_path)
        self.assertTrue(os.path.isfile(out_path))
        self.assertEqual(_hash_file(source), before_hash)

        result = read_docx(out_path)
        self.assertEqual(result.paragraphs[0].text, HEADING_TEXT)
        self.assertEqual(result.paragraphs[1].text, replacement)
        self.assertEqual(result.paragraphs[2].text, BODY2)
        self.assertEqual(result.paragraphs[4].text, REF_TEXT)
        self.assertEqual(len(result.paragraphs), 5)

    def test_write_docx_ignores_unknown_index(self) -> None:
        source = os.path.join(self.tmp, "source.docx")
        _make_sample_docx(source)
        out_path = os.path.join(self.tmp, "unknown.docx")
        write_docx(source, {99: "不应写入任何段落"}, out_path)
        result = read_docx(out_path)
        self.assertEqual(result.paragraphs[1].text, BODY1)
        self.assertNotIn("不应写入任何段落", [p.text for p in result.paragraphs])

    def test_write_docx_preserves_first_run_format(self) -> None:
        source = os.path.join(self.tmp, "runs.docx")
        document = docx.Document()
        paragraph = document.add_paragraph()
        first = paragraph.add_run("旧的段落文本内容")
        first.bold = True
        first.font.size = Pt(12)
        second = paragraph.add_run("尾巴文字")
        second.italic = True
        document.save(source)

        out_path = os.path.join(self.tmp, "runs_out.docx")
        write_docx(source, {0: "新的段落文本"}, out_path)

        result_doc = docx.Document(out_path)
        result_paragraph = result_doc.paragraphs[0]
        runs = result_paragraph.runs
        self.assertEqual(result_paragraph.text, "新的段落文本")
        self.assertEqual(len(runs), 2)
        self.assertEqual(runs[0].text, "新的段落文本")
        self.assertTrue(runs[0].bold)
        self.assertEqual(runs[0].font.size, Pt(12))
        self.assertFalse(runs[1].text)
        self.assertTrue(runs[1].italic)

    def test_write_docx_rejects_same_path(self) -> None:
        source = os.path.join(self.tmp, "same.docx")
        _make_sample_docx(source)
        with self.assertRaises(ValueError) as ctx:
            write_docx(source, {0: "覆盖"}, source)
        self.assertIn("源文件", str(ctx.exception))


class ReportHtmlTest(_TempDirMixin, unittest.TestCase):
    def _render(self, report: ScoreReport, rewrite=None) -> str:
        out_path = os.path.join(self.tmp, "html", "report.html")
        returned = render_html(report, out_path, rewrite)
        self.assertEqual(returned, out_path)
        self.assertTrue(os.path.isfile(out_path))
        with open(out_path, "r", encoding="utf-8") as handle:
            return handle.read()

    def test_html_contains_disclaimer_and_basic_blocks(self) -> None:
        html = self._render(_sample_report())
        self.assertIn(DISCLAIMER_ZH, html)
        self.assertIn("AI 痕迹指数", html)
        self.assertIn("参与评分段落", html)
        self.assertIn("正文总字数", html)
        self.assertIn("高风险段数", html)
        self.assertIn("最高风险段落", html)
        self.assertIn("段落风险排行", html)
        self.assertIn("生成时间", html)
        self.assertIn(VERSION, html)
        self.assertIn("未校准（uncalibrated）", html)

    def test_html_disclaimer_visible_and_not_collapsible(self) -> None:
        html = self._render(_sample_report())
        start = html.index('id="disclaimer"')
        end = html.index("</section>", start)
        section = html[start:end]
        self.assertIn("font-size:14.5px", section)
        self.assertNotIn("details", section)
        self.assertIn("重要说明（必读）", section)
        self.assertIn("不是知网 AIGC 率", section)

    def test_html_escapes_user_text(self) -> None:
        html = self._render(_sample_report())
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("&amp;", html)
        self.assertIn("&lt;b&gt;并拆句&lt;/b&gt;", html)

    def test_html_has_no_external_resources(self) -> None:
        html = self._render(_sample_report())
        lowered = html.lower()
        self.assertNotIn("http://", lowered)
        self.assertNotIn("https://", lowered)
        self.assertNotIn("<link", lowered)
        self.assertNotIn("<img", lowered)
        self.assertNotIn("src=", lowered)

    def test_html_sorts_paragraphs_by_risk(self) -> None:
        html = self._render(_sample_report())
        self.assertLess(html.index("&lt;script&gt;"), html.index(LOW_TEXT))

    def test_html_rewrite_comparison(self) -> None:
        report = _sample_report()
        rewrite = {
            0: RewriteResult(
                index=0,
                original="原始 & 句子",
                revised="改写后 <更自然>",
                score_before=72.5,
                score_after=31.2,
                attempts=2,
                ok=True,
            ),
            1: RewriteResult(
                index=1,
                original="另一段原文",
                revised="",
                score_before=80.0,
                score_after=80.0,
                attempts=3,
                ok=False,
                error="改写无变化",
            ),
        }
        html = self._render(report, rewrite)
        self.assertIn("改写前后对照", html)
        self.assertIn("72.5", html)
        self.assertIn("31.2", html)
        self.assertIn("改写后 &lt;更自然&gt;", html)
        self.assertIn("改写无变化", html)
        self.assertIn("未完成", html)

    def test_html_without_rewrite_has_no_rewrite_section(self) -> None:
        html = self._render(_sample_report())
        self.assertNotIn("改写前后对照", html)

    def test_html_empty_report_renders(self) -> None:
        report = ScoreReport(
            document=Document(path="空.docx", title="空文档"),
            paragraphs=[],
            overall=0.0,
            level="低",
            level_counts={},
            meta={"engine": "stylometric-heuristic"},
        )
        html = self._render(report)
        self.assertIn("没有参与评分的段落", html)
        self.assertIn("—", html)
        self.assertIn(DISCLAIMER_ZH, html)


class ReportMarkdownTest(_TempDirMixin, unittest.TestCase):
    def test_markdown_structure_and_escaping(self) -> None:
        report = _sample_report()
        report.paragraphs[0].suggestion = "拆长句 | 去套话"
        markdown = render_markdown(report)

        self.assertIn(DISCLAIMER_ZH, markdown)
        self.assertIn("| 指标 | 数值 |", markdown)
        self.assertIn("## 段落风险排行", markdown)
        self.assertIn("拆长句 \\| 去套话", markdown)
        self.assertIn(VERSION, markdown)
        self.assertTrue(markdown.endswith("\n"))

        table_start = markdown.index("| 序号 | 分数 | 等级 | Top 信号 | 一句话建议 |")
        self.assertLess(
            markdown.index("88.0", table_start),
            markdown.index("22.5", table_start),
        )

    def test_markdown_rewrite_section(self) -> None:
        report = _sample_report()
        rewrite = {
            0: RewriteResult(
                index=0,
                original="原始文本",
                revised="改写后的文本",
                score_before=70.0,
                score_after=30.0,
                attempts=1,
                ok=True,
            )
        }
        markdown = render_markdown(report, rewrite)
        self.assertIn("## 改写前后对照", markdown)
        self.assertIn("70.0 → 30.0", markdown)
        self.assertIn("改写后的文本", markdown)

    def test_markdown_empty_report(self) -> None:
        report = ScoreReport(
            document=Document(path="空.docx", title="空文档"),
            paragraphs=[],
            overall=0.0,
            level="低",
        )
        markdown = render_markdown(report)
        self.assertIn("（没有参与评分的段落。）", markdown)
        self.assertIn("—", markdown)


if __name__ == "__main__":
    unittest.main()
