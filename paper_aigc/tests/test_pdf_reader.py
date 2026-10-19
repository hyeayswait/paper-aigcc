"""T4 单元测试：PDF 解析 / 统一文档加载。

临时 PDF 用 pymupdf 现场生成，全部写到 tempfile.mkdtemp() 下，不往仓库丢文件；
不依赖网络与 API Key，也不依赖 T1/T2 的评分逻辑。
"""
from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from unittest import mock

import docx

from paper_aigc.io import pdf_reader
from paper_aigc.io.document_loader import read_document, supported_extensions
from paper_aigc.io.pdf_reader import is_pdf_supported, read_pdf

fitz = pdf_reader.fitz

TITLE = "第一章 绪论"
BODY1 = "本研究围绕论文的文体特征展开分析，并给出可复核的指标定义。"
BODY2 = "在方法上，本文采用分段统计与加权汇总的方式构造指数。"
BODY3 = "最后，本文给出实验设计与结果分析，并讨论方法的适用边界。"
HEADER = "某大学硕士学位论文"
REF_TEXT = "[1] 张三. 论文题名[J]. 期刊名称, 2020."
SHORT_TEXT = "结论如下。"


class _TempDirMixin:
    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="paperaigc_pdf_")

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)


def _insert(page, x: float, y: float, text: str, size: int = 12) -> None:
    page.insert_text((x, y), text, fontname="china-s", fontsize=size)


def _make_sample_pdf(path: str) -> None:
    """3 页 PDF：第 1 页标题 + 正文，第 2/3 页带重复页眉与页码。"""
    document = fitz.open()
    page1 = document.new_page()
    _insert(page1, 72, 90, TITLE, 16)
    _insert(page1, 72, 140, BODY1)

    page2 = document.new_page()
    _insert(page2, 72, 40, HEADER, 9)
    _insert(page2, 72, 120, BODY2)
    _insert(page2, 300, 780, "2", 9)

    page3 = document.new_page()
    _insert(page3, 72, 40, HEADER, 9)
    _insert(page3, 72, 120, BODY3)
    _insert(page3, 300, 780, "3", 9)

    document.save(path)
    document.close()


@unittest.skipUnless(fitz is not None, "pymupdf 不可用，跳过 PDF 生成相关测试")
class PdfReaderTest(_TempDirMixin, unittest.TestCase):
    def test_is_pdf_supported_true(self) -> None:
        self.assertTrue(is_pdf_supported())

    def test_read_pdf_structure(self) -> None:
        path = os.path.join(self.tmp, "sample.pdf")
        _make_sample_pdf(path)
        document = read_pdf(path)

        self.assertEqual(document.path, path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual(
            [p.text for p in document.paragraphs], [TITLE, BODY1, BODY2, BODY3]
        )
        self.assertEqual([p.index for p in document.paragraphs], [0, 1, 2, 3])
        self.assertTrue(document.paragraphs[0].is_heading)
        self.assertFalse(document.paragraphs[1].is_heading)
        self.assertFalse(document.paragraphs[1].is_reference)
        self.assertEqual(document.body_indices, [1, 2, 3])
        self.assertEqual(
            document.total_words,
            sum(document.paragraphs[i].char_count for i in document.body_indices),
        )

    def test_read_pdf_removes_header_and_page_numbers(self) -> None:
        path = os.path.join(self.tmp, "header.pdf")
        _make_sample_pdf(path)
        document = read_pdf(path)
        texts = [p.text for p in document.paragraphs]
        self.assertNotIn(HEADER, texts)
        self.assertNotIn("2", texts)
        self.assertNotIn("3", texts)

    def test_read_pdf_excludes_reference_and_short_lines(self) -> None:
        path = os.path.join(self.tmp, "ref.pdf")
        document = fitz.open()
        page = document.new_page()
        _insert(page, 72, 90, TITLE, 16)
        _insert(page, 72, 140, BODY1)
        _insert(page, 72, 180, REF_TEXT)
        _insert(page, 72, 220, SHORT_TEXT)
        document.save(path)
        document.close()

        read = read_pdf(path)
        self.assertEqual(
            [p.text for p in read.paragraphs], [TITLE, BODY1, REF_TEXT, SHORT_TEXT]
        )
        self.assertTrue(read.paragraphs[2].is_reference)
        self.assertLess(read.paragraphs[3].char_count, 10)
        self.assertEqual(read.body_indices, [1])

    def test_read_pdf_title_falls_back_to_file_name(self) -> None:
        path = os.path.join(self.tmp, "no_heading.pdf")
        document = fitz.open()
        page = document.new_page()
        _insert(page, 72, 120, BODY1)
        document.save(path)
        document.close()

        read = read_pdf(path)
        self.assertEqual(read.title, "no_heading")
        self.assertEqual(read.body_indices, [0])

    def test_read_pdf_blank_page_returns_empty_document(self) -> None:
        path = os.path.join(self.tmp, "blank.pdf")
        document = fitz.open()
        document.new_page()
        document.save(path)
        document.close()

        read = read_pdf(path)
        self.assertEqual(read.paragraphs, [])
        self.assertEqual(read.body_indices, [])
        self.assertEqual(read.total_words, 0)
        self.assertEqual(read.title, "blank")

    def test_read_pdf_missing_file_raises_value_error(self) -> None:
        missing = os.path.join(self.tmp, "not_here.pdf")
        with self.assertRaises(ValueError) as ctx:
            read_pdf(missing)
        self.assertIn("不存在", str(ctx.exception))

    def test_read_pdf_wrong_extension_raises_value_error(self) -> None:
        wrong = os.path.join(self.tmp, "paper.txt")
        with open(wrong, "w", encoding="utf-8") as handle:
            handle.write("hello")
        with self.assertRaises(ValueError) as ctx:
            read_pdf(wrong)
        self.assertIn(".pdf", str(ctx.exception))

    def test_read_pdf_broken_file_raises_value_error(self) -> None:
        broken = os.path.join(self.tmp, "broken.pdf")
        with open(broken, "wb") as handle:
            handle.write(b"this is not a real pdf")
        with self.assertRaises(ValueError) as ctx:
            read_pdf(broken)
        self.assertIn("无法解析", str(ctx.exception))


class PdfDegradeTest(_TempDirMixin, unittest.TestCase):
    def test_without_pymupdf_raises_value_error(self) -> None:
        path = os.path.join(self.tmp, "dummy.pdf")
        with open(path, "wb") as handle:
            handle.write(b"%PDF-1.4\n")
        with mock.patch.object(pdf_reader, "fitz", None):
            self.assertFalse(pdf_reader.is_pdf_supported())
            with self.assertRaises(ValueError) as ctx:
                pdf_reader.read_pdf(path)
        self.assertIn("pymupdf", str(ctx.exception))


class DocumentLoaderTest(_TempDirMixin, unittest.TestCase):
    def test_supported_extensions(self) -> None:
        extensions = supported_extensions()
        self.assertIsInstance(extensions, list)
        for ext in (".docx", ".pdf", ".txt", ".md"):
            self.assertIn(ext, extensions)
        extensions.append(".xyz")
        self.assertNotIn(".xyz", supported_extensions())

    def test_read_document_txt(self) -> None:
        path = os.path.join(self.tmp, "sample.txt")
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(TITLE + "\n\n" + BODY1 + "\n\n" + BODY2 + "\n")
        document = read_document(path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual([p.text for p in document.paragraphs], [TITLE, BODY1, BODY2])
        self.assertTrue(document.paragraphs[0].is_heading)
        self.assertEqual(document.body_indices, [1, 2])

    def test_read_document_txt_gb18030(self) -> None:
        path = os.path.join(self.tmp, "gbk.txt")
        with open(path, "w", encoding="gb18030") as handle:
            handle.write(TITLE + "\n\n" + BODY1 + "\n")
        document = read_document(path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual(document.body_indices, [1])

    def test_read_document_markdown_headings(self) -> None:
        path = os.path.join(self.tmp, "sample.md")
        with open(path, "w", encoding="utf-8") as handle:
            handle.write("# " + TITLE + "\n\n" + BODY1 + "\n\n## 研究背景\n\n" + BODY2 + "\n")
        document = read_document(path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual(
            [p.text for p in document.paragraphs],
            [TITLE, BODY1, "研究背景", BODY2],
        )
        self.assertEqual([p.is_heading for p in document.paragraphs], [True, False, True, False])
        self.assertEqual(document.body_indices, [1, 3])

    @unittest.skipUnless(fitz is not None, "pymupdf 不可用，跳过 PDF 分派测试")
    def test_read_document_dispatches_pdf(self) -> None:
        path = os.path.join(self.tmp, "sample.pdf")
        _make_sample_pdf(path)
        document = read_document(path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual(document.body_indices, [1, 2, 3])

    def test_read_document_dispatches_docx(self) -> None:
        path = os.path.join(self.tmp, "sample.docx")
        source = docx.Document()
        heading = source.add_paragraph(TITLE)
        heading.style = source.styles["Heading 1"]
        source.add_paragraph(BODY1)
        source.save(path)

        document = read_document(path)
        self.assertEqual(document.title, TITLE)
        self.assertEqual([p.text for p in document.paragraphs], [TITLE, BODY1])
        self.assertEqual(document.body_indices, [1])

    def test_read_document_unknown_extension_raises_value_error(self) -> None:
        path = os.path.join(self.tmp, "paper.xyz")
        with open(path, "w", encoding="utf-8") as handle:
            handle.write("hello")
        with self.assertRaises(ValueError) as ctx:
            read_document(path)
        self.assertIn("不支持", str(ctx.exception))

    def test_read_document_missing_file_raises_value_error(self) -> None:
        with self.assertRaises(ValueError) as ctx:
            read_document(os.path.join(self.tmp, "missing.txt"))
        self.assertIn("不存在", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
