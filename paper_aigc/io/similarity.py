# -*- coding: utf-8 -*-
"""本地库查重：用「连续 N 字片段」重合度衡量段落与本地参考文档的重复情况。

重要边界：
- 只与 corpus_dir 里的本地文档比对，不联网、不包含学术库/互联网；
- 因此结果不能等同于 PaperPass / 知网 / 万方的查重率；
- 用途：查是否抄了参考文献、查自我重复、查与同门往届论文的重合。
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

import docx

from ..core.textnorm import normalize_text
from ..core.types import Document

DEFAULT_NGRAM = 13
_SUPPORTED = (".txt", ".md", ".docx")

_CJK_RE = re.compile(r"[\u4e00-\u9fff]")
_KEEP_RE = re.compile(r"[\u4e00-\u9fffA-Za-z0-9]")


@dataclass
class MatchSource:
    source: str              # 文件名
    matched_grams: int       # 命中的片段数
    ratio: float             # 占该段片段的比例
    snippet: str = ""        # 示例重合内容


@dataclass
class ParagraphSimilarity:
    index: int
    text: str
    ratio: float
    sources: List[MatchSource] = field(default_factory=list)


@dataclass
class SimilarityReport:
    overall: float = 0.0                       # 按字符加权的总复制比（0-1）
    paragraphs: List[ParagraphSimilarity] = field(default_factory=list)
    internal_duplicates: List[Tuple[int, int, float]] = field(default_factory=list)
    corpus_files: int = 0
    corpus_chars: int = 0
    ngram: int = DEFAULT_NGRAM
    note: str = ""
    self_compare: List[Tuple[str, float]] = field(default_factory=list)

    @property
    def overall_percent(self) -> float:
        return round(self.overall * 100.0, 2)

    def sorted_by_risk(self) -> List[ParagraphSimilarity]:
        return sorted(self.paragraphs, key=lambda item: item.ratio, reverse=True)


def _normalize(text: str) -> str:
    """只保留汉字/字母/数字，去掉标点与空白，避免标点差异影响比对。"""
    return "".join(_KEEP_RE.findall(text or "")).lower()


def _grams(text: str, ngram: int) -> List[str]:
    if ngram <= 0:
        return []
    seq = _normalize(text)
    if len(seq) < ngram:
        return []
    return [seq[i:i + ngram] for i in range(len(seq) - ngram + 1)]


def load_corpus(corpus_dir: str) -> List[Tuple[str, str]]:
    """读取本地比对库，返回 [(文件名, 正文)]。recursive 扫描 .txt/.md/.docx。"""
    items: List[Tuple[str, str]] = []
    if not corpus_dir or not os.path.isdir(corpus_dir):
        return items
    for root, _dirs, files in os.walk(corpus_dir):
        for name in sorted(files):
            if not name.lower().endswith(_SUPPORTED):
                continue
            path = os.path.join(root, name)
            try:
                if name.lower().endswith(".docx"):
                    document = docx.Document(path)
                    text = "\n".join(p.text for p in document.paragraphs)
                else:
                    with open(path, encoding="utf-8", errors="ignore") as handle:
                        text = handle.read()
            except Exception:
                continue
            if text.strip():
                items.append((os.path.relpath(path, corpus_dir), normalize_text(text)))
    return items


def to_meta(report: SimilarityReport, top: int = 15) -> dict:
    """序列化成可放进 ScoreReport.meta 的 dict，供报告渲染。"""
    return {
        "overall_percent": report.overall_percent,
        "corpus_files": report.corpus_files,
        "corpus_chars": report.corpus_chars,
        "ngram": report.ngram,
        "note": report.note,
        "paragraphs": [
            {
                "index": item.index,
                "display": item.index + 1,
                "ratio_percent": round(item.ratio * 100.0, 1),
                "sources": [
                    {"source": src.source, "ratio_percent": round(src.ratio * 100.0, 1),
                     "snippet": src.snippet}
                    for src in item.sources
                ],
            }
            for item in report.sorted_by_risk()[:top]
            if item.ratio > 0
        ],
        "internal_duplicates": [
            {"a": a + 1, "b": b + 1, "overlap_percent": round(overlap * 100.0, 1)}
            for a, b, overlap in report.internal_duplicates
        ],
    }


def check_similarity(doc: Document, corpus_dir: str, ngram: int = DEFAULT_NGRAM) -> SimilarityReport:
    """对 Document 的正文段落做本地库查重。"""
    report = SimilarityReport(ngram=ngram)
    corpus = load_corpus(corpus_dir)
    report.corpus_files = len(corpus)
    report.corpus_chars = sum(len(text) for _name, text in corpus)

    if not corpus:
        report.note = ("未找到本地比对库：请把参考文档（.txt/.md/.docx）放进 %s。"
                       "本功能不联网，只与本地库比对。" % (corpus_dir or "corpus 目录"))
        return report

    # ngram -> 出现的来源集合
    index: Dict[str, set] = {}
    for position, (_name, text) in enumerate(corpus):
        for gram in _grams(text, ngram):
            index.setdefault(gram, set()).add(position)

    total_chars = 0
    matched_chars = 0
    paragraphs: List[ParagraphSimilarity] = []
    for index_number in doc.body_indices:
        paragraph = next((p for p in doc.paragraphs if p.index == index_number), None)
        if paragraph is None:
            continue
        grams = _grams(paragraph.text, ngram)
        if not grams:
            continue
        hit_by_source: Dict[int, List[int]] = {}
        matched_positions: List[int] = []
        for offset, gram in enumerate(grams):
            sources = index.get(gram)
            if not sources:
                continue
            matched_positions.append(offset)
            for position in sources:
                hit_by_source.setdefault(position, []).append(offset)
        ratio = len(matched_positions) / float(len(grams))
        sources_info: List[MatchSource] = []
        for position, offsets in sorted(hit_by_source.items(), key=lambda kv: -len(kv[1])):
            offsets.sort()
            start = offsets[0]
            snippet = "".join(_normalize(paragraph.text)[start:start + min(60, ngram * 3)])
            sources_info.append(MatchSource(
                source=corpus[position][0],
                matched_grams=len(offsets),
                ratio=round(len(offsets) / float(len(grams)), 4),
                snippet=snippet,
            ))
        sources_info.sort(key=lambda item: -item.matched_grams)
        paragraphs.append(ParagraphSimilarity(
            index=index_number, text=paragraph.text, ratio=round(ratio, 4),
            sources=sources_info[:3],
        ))
        total_chars += len(paragraph.text)
        matched_chars += int(round(ratio * len(paragraph.text)))

    report.paragraphs = paragraphs
    report.overall = round(matched_chars / float(total_chars), 4) if total_chars else 0.0

    # 自我重复：段落之间互相高度重合
    for i in range(len(paragraphs)):
        for j in range(i + 1, len(paragraphs)):
            a = _grams(paragraphs[i].text, ngram)
            b = _grams(paragraphs[j].text, ngram)
            if not a or not b:
                continue
            overlap = len(set(a) & set(b)) / float(min(len(set(a)), len(set(b))))
            if overlap >= 0.5:
                report.internal_duplicates.append(
                    (paragraphs[i].index, paragraphs[j].index, round(overlap, 3)))

    # 自比保险：比对库里若有文件与本文件高度重合，说明可能误把本论文放进了库
    doc_grams = set(_grams("".join(item.text for item in paragraphs), ngram))
    self_compare = []
    if doc_grams:
        for name, corpus_text in corpus:
            corpus_grams = set(_grams(corpus_text, ngram))
            if not corpus_grams:
                continue
            overlap = len(doc_grams & corpus_grams) / float(len(doc_grams))
            if overlap >= 0.5:
                self_compare.append((name, overlap))
    report.self_compare = self_compare

    report.note = ("本地库查重：与 %d 个本地文件（%d 字）比对，%d 字连续片段为重合判据。"
                   "不联网、不含学术库/互联网，数字不等同于 PaperPass/知网。" % (
                       report.corpus_files, report.corpus_chars, ngram))
    if self_compare:
        report.note += (" ⚠ 比对库中有文件与本文件高度重合（%s），"
                        "可能是把本论文本身放进了比对库，结果含自比成分，请移出后重跑。"
                        % "、".join("%s %.0f%%" % (name, ratio * 100) for name, ratio in self_compare[:3]))
    return report
