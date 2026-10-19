# -*- coding: utf-8 -*-
"""生成一份带「格式硬伤」的样例，用于演示规范自检。

    python samples/make_lint_sample.py
"""
from __future__ import annotations

import os

from docx import Document


def main() -> str:
    root = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(root, "问题样例_格式硬伤.docx")
    doc = Document()

    h = doc.add_paragraph()
    h.style = doc.styles["Heading 1"]
    h.add_run("一、研究方法")

    doc.add_paragraph("本文使用,问卷调查与访谈相结合的方法.样本覆盖三所高校,共回收问卷两百份.")
    doc.add_paragraph(
        "这种方法的好处是能够在较短时间内获得较大范围的样本，但也存在明显的局限，"
        "比如受访者对问题的理解可能不一致、问卷回收率难以保证，而且在分析过程中如何处理缺失值、"
        "如何界定核心概念的边界以及如何解释指标之间的相关性，都需要在方法部分交代清楚否则读者很难判断结论是否可靠。"
    )
    doc.add_paragraph("相关数据见表3与图2，具体口径见表1。")

    cap1 = doc.add_paragraph("表1 样本分布")
    cap2 = doc.add_paragraph("表2 变量定义")

    doc.add_paragraph("本研究中净资产收益率是最核心的指标，同时我们也参考了同行业 ROE 的均值水平。")

    ref_head = doc.add_paragraph()
    ref_head.style = doc.styles["Heading 1"]
    ref_head.add_run("参考文献")
    doc.add_paragraph("[1] 张三. 某文献[J].")
    doc.add_paragraph("李四. 某专著[M]. 北京: 某出版社, 2020.")

    doc.save(out)
    print("saved:", out)
    return out


if __name__ == "__main__":
    main()
